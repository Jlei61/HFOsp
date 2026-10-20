#!/usr/bin/env python3
"""Apply the 2026-10-19 Fig1 comments in a separate, reviewable A–G package.

Reuse accepted PDF layers, the HFO painter and the original MI painter. Only
waveform display units, labels, panel boundaries and MI axis size are changed.
Requires numpy, matplotlib, Pillow, pypdf, python-pptx and pdftoppm.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

os.environ.setdefault("MPLCONFIGDIR", "/tmp/fig1-author-comments-mpl")
sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

# The accepted release keeps historical layout helpers inside its frozen code
# package. Resolve missing helpers there so a clean checkout can reproduce this
# revision without the author's untracked development scripts.
FROZEN_FIGURE_CODE = (ROOT / "results/paper-ready-figure/fig1/source/"
                     "code_snapshot/scripts/paper_figures")
import scripts.paper_figures as figure_scripts
figure_scripts.__path__.append(str(FROZEN_FIGURE_CODE))
sys.path.append(str(FROZEN_FIGURE_CODE))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import ScalarFormatter
from matplotlib.transforms import ScaledTranslation
import numpy as np
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject, RectangleObject

from scripts.paper_figures import build_fig1_three_column_review as layout
from scripts.paper_figures import build_fig1_readability_review as shared
from scripts.paper_figures.build_fig1_current import verify
from scripts.paper_figures.build_fig1_y1_local_rank import style_axis, TYPE
from scripts.paper_figures.revise_fig1_lower_row_gap import clipped_page

BASE = ROOT / "results/paper-ready-figure/fig1"
VERSION = "author_comments_split_bc_20261019"
WIDTH, HEIGHT = 19.5, 15.65
DPI = 300
# Non-overlapping deletion rectangles, in bottom-up canvas inches. The
# original C/D/E/F letters are replaced by D/E/F/G. C's spectra remain intact.
REMOVE = [
    (10.60, 9.95, 14.02, 15.65),
    (14.03, 4.96, 19.46, 9.94),
    (0.00, 9.73, 0.50, 10.28),
    (14.03, 9.94, 14.70, 10.28),
    (0.00, 4.96, 0.50, 5.48),
]
LETTERS = [("B", 10.72, 15.55), ("C", 14.18, 15.55),
           ("D", .08, 10.17), ("E", 14.18, 10.17),
           ("F", .08, 5.39), ("G", 14.18, 5.39)]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def rect(x, y, w, h):
    return [x/WIDTH, y/HEIGHT, w/WIDTH, h/HEIGHT]


def render(pdf, png, dpi=DPI):
    subprocess.run(["pdftoppm", "-singlefile", "-r", str(dpi), "-png",
                    str(pdf), str(png.with_suffix(""))], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)


def save_pdf(page, path):
    writer = PdfWriter()
    writer.add_page(page)
    with path.open("wb") as stream:
        writer.write(stream)


def crop_pdf(source, box_inches, destination):
    page = PdfReader(source).pages[0]
    box = [v*72 for v in box_inches]
    clipped = clipped_page(page, [box])
    clipped.mediabox = RectangleObject(box)
    clipped.cropbox = RectangleObject(box)
    save_pdf(clipped, destination)


def preserved_page(page):
    """Join E's deletion regions into one polygon to avoid a clipping seam."""
    result = copy.copy(page)
    paths = [f"0 0 {WIDTH*72} {HEIGHT*72} re"]
    for i in (0, 2, 4):
        x0, y0, x1, y1 = (v*72 for v in REMOVE[i])
        paths.append(f"{x0} {y0} {x1-x0} {y1-y0} re")
    vertices = [(14.03, 4.96), (19.46, 4.96), (19.46, 9.94),
                (14.70, 9.94), (14.70, 10.28), (14.03, 10.28)]
    paths.extend(f"{x*72} {y*72} {'m' if i == 0 else 'l'}"
                 for i, (x, y) in enumerate(vertices))
    paths.extend(["h", "W* n"])
    stream = DecodedStreamObject()
    stream.set_data(b"q\n" + "\n".join(paths).encode() + b"\n" +
                    page.get_contents().get_data() + b"\nQ\n")
    result[NameObject("/Contents")] = stream
    return result


def build_panels(output, accepted):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": TYPE.annotation,
                         "pdf.fonttype": 42, "svg.fonttype": "none",
                         "axes.unicode_minus": False})
    fig = plt.figure(figsize=(WIDTH, HEIGHT))
    old_output = layout.OUT
    try:
        layout.OUT = output
        axes, hfo_metadata = layout.draw_hfo(
            fig, 10.75, showcase_path=BASE / "data/b/hfo_showcase.npz")
    finally:
        layout.OUT = old_output
    for i, ax in enumerate(axes):
        ax.set_position(rect(11.83, 10.69+(2-i)*1.47, 1.87, 1.22))
        style_axis(ax)
        ax.set_title("")
        ax.set(xlim=(0, .6), xticks=[0, .3, .6])
        ax.set_xticklabels(["0", "0.3", "0.6"])
        ax.tick_params(axis="x", labelbottom=i == 2)
        if i:
            ax.set(ylim=(0, 240), yticks=[0, 120, 240])
            ax.set_ylabel("Freq (Hz)")
    wave = axes[0]
    original_lines = [line.get_ydata().copy() for line in wave.lines]
    for line, volts in zip(wave.lines, original_lines):
        line.set_ydata(volts*1e6)
    wave.set(ylim=(-100, 150), yticks=[-100, 0, 100], ylabel="LFP (μV)")
    formatter = ScalarFormatter(useOffset=False)
    formatter.set_scientific(False)
    wave.yaxis.set_major_formatter(formatter)
    for ax in axes:
        ax.yaxis.set_label_coords(-.43, .5)
    axes[-1].set_xlabel("Time (s)")

    with np.load(BASE / "data/b/hfo_showcase.npz") as payload:
        assert len(wave.lines) == 179
        for line, original in zip(wave.lines[:-1], payload["snippets_V"]):
            np.testing.assert_array_equal(line.get_ydata(), original*1e6)
        np.testing.assert_array_equal(wave.lines[-1].get_ydata(),
                                      payload["snippets_V"].mean(0)*1e6)
        for axis, key in zip(axes[1:], ("raw_spec", "normalized_spec")):
            np.testing.assert_array_equal(axis.collections[0].get_array(), payload[key])
    with np.load(BASE / "data/b/hfo_showcase.npz") as original, \
            np.load(output / "source/hfo_showcase.npz") as saved:
        for key in original.files:
            np.testing.assert_array_equal(original[key], saved[key])

    records = json.loads((BASE / "data/df/cohort_plot_records.json").read_text())
    shared.old._assert_masked_mi_records(records)
    start = len(fig.axes)
    mi = shared.draw_d(fig, rect(14, 5, 5, 5), records, None, show_mechanism=False)
    stats = fig.axes[start]
    stats.set_position(rect(14.98, 5.82, 3.94, 3.30))
    style_axis(stats)
    for text in stats.texts:
        if text.get_text() in ("Yuquan", "Epilepsiae"):
            text.set_y(0)
            text.set_transform(stats.get_xaxis_transform() +
                               ScaledTranslation(0, -.50, fig.dpi_scale_trans))
            text.set_va("top")
    assert mi == accepted["summaries"]["D"]
    # Show the same 80 observed/null subject values at their exact y positions.
    from matplotlib.collections import PathCollection
    points = [c for c in stats.collections if isinstance(c, PathCollection)]
    for i, dataset in enumerate(("yuquan", "epilepsiae")):
        rows = [r for r in records if r["dataset"] == dataset]
        for offset, key in enumerate(("mi_mean", "permuted_mean_median")):
            np.testing.assert_array_equal(points[2*i+offset].get_offsets()[:, 1],
                                          [r["legacy_mi"][key] for r in rows])
    fig.canvas.draw()
    boxes = {"B": layout.bounds(fig, axes).padded(.04),
             "E": layout.bounds(fig, [stats]).padded(.04)}
    assert boxes["B"].x0 > 10.60 and boxes["B"].x1 < 14.02
    assert boxes["E"].y0 > 5.02 and boxes["E"].y1 < 9.94
    assert all(not ax.get_title() for ax in axes)
    for key, box in boxes.items():
        pdf = output / f"figures/fig1-panel{key.lower()}.pdf"
        fig.savefig(pdf, dpi=DPI, facecolor="white", bbox_inches=box)
        render(pdf, pdf.with_suffix(".png"))

    for letter, x, y in LETTERS:
        fig.text(x/WIDTH, y/HEIGHT, letter, fontsize=TYPE.panel_letter,
                 fontweight="bold", va="top")
    overlay = output / "source/revision_overlay.pdf"
    fig.savefig(overlay, dpi=DPI, transparent=True)
    plt.close(fig)
    return overlay, dict(
        waveform_source_unit="V", waveform_display_unit="μV", waveform_scale=1e6,
        waveform_traces=178, waveform_mean_trace=True,
        source_arrays_exact=True, spectra_exact=True,
        B_time_limits_s=[0, .6], B_time_ticks_s=[0, .3, .6],
        B_titles_removed=["Raw", "Normalized", "HFO n=178"],
        E_statistical_summary_exact=True, E_all_80_subject_values_exact=True,
        E_axis_inches=[14.98, 5.82, 3.94, 3.30],
        new_panel_bounds_inches={k: v.extents.tolist() for k, v in boxes.items()},
        hfo_source=hfo_metadata, E_summary=mi)


def make_pptx(output):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    names = ["complete-layout"] + [f"panel{x}" for x in "abcdefg"]
    for name in names:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        path = output / f"figures/fig1-{name}.png"
        with Image.open(path) as picture:
            w, h = picture.size
        scale = min(WIDTH/w, HEIGHT/h)
        slide.shapes.add_picture(str(path), Inches((WIDTH-w*scale)/2),
                                 Inches((HEIGHT-h*scale)/2),
                                 width=Inches(w*scale), height=Inches(h*scale))
    destination = output / "figures/fig1-author-comments.pptx"
    prs.save(destination)
    assert len(Presentation(destination).slides) == 8


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=BASE / "candidates" / VERSION)
    args = parser.parse_args()
    output = args.output_dir.resolve()
    if output == BASE.resolve() or BASE.resolve() not in output.parents:
        raise ValueError("Write the revision within a separate Fig1 candidate directory.")
    output.mkdir(parents=True, exist_ok=True)
    (output / "figures").mkdir(exist_ok=True)
    (output / "source").mkdir(exist_ok=True)
    accepted_report = verify(BASE)
    pointer_hash = sha(BASE / "current_revision.json")
    accepted = json.loads((BASE / "metadata.json").read_text())
    overlay, checks = build_panels(output, accepted)
    print("B/E redrawn from frozen accepted numerical inputs", flush=True)

    base_pdf = BASE / "figures/fig1-complete-layout.pdf"
    page = PdfReader(base_pdf).pages[0]
    clipped = preserved_page(page)
    clipped.merge_page(PdfReader(overlay).pages[0])
    complete_pdf = output / "figures/fig1-complete-layout.pdf"
    save_pdf(clipped, complete_pdf)
    render(complete_pdf, complete_pdf.with_suffix(".png"))

    # Independent preserved panels retain the accepted bytes under new letters.
    mapping = {"a": "a", "c": "d", "e": "f", "f": "g"}
    for old, new in mapping.items():
        for ext in ("pdf", "png"):
            source = BASE / f"figures/fig1-panel{old}.{ext}"
            dest = output / f"figures/fig1-panel{new}.{ext}"
            shutil.copy2(source, dest)
            assert sha(source) == sha(dest)
    c_box = (14.60, 9.95, 19.40, 15.22)
    c_pdf = output / "figures/fig1-panelc.pdf"
    crop_pdf(base_pdf, c_box, c_pdf)
    render(c_pdf, c_pdf.with_suffix(".png"))
    for ext in ("pdf", "png"):
        name = f"fig1-spectrum-full-window-check.{ext}"
        shutil.copy2(BASE / "source/layout_input/figures" / name, output / "source" / name)
    shutil.copy2(BASE / "spectrum_contract.json", output / "spectrum_contract.json")
    shutil.copy2(BASE / "data/df/cohort_plot_records.json", output / "source/cohort_plot_records.json")

    # Compare using the same PDF renderer: all regions outside the requested
    # changes, including C's spectra/centroids and A/D/F/G data, must be exact.
    before_png = output / "source/accepted_pdf_render.png"
    render(base_pdf, before_png)
    before = np.asarray(Image.open(before_png).convert("RGB"))
    after = np.asarray(Image.open(complete_pdf.with_suffix(".png")).convert("RGB"))
    assert before.shape == after.shape
    allowed = np.zeros(before.shape[:2], dtype=bool)
    change_boxes = REMOVE + [(14.08, 15.05, 14.62, 15.65)]
    for x0, y0, x1, y1 in change_boxes:
        left, right = max(0, int(x0*DPI)-3), min(before.shape[1], int(x1*DPI)+4)
        top, bottom = max(0, int((HEIGHT-y1)*DPI)-3), min(before.shape[0], int((HEIGHT-y0)*DPI)+4)
        allowed[top:bottom, left:right] = True
    changed = np.any(before != after, axis=2)
    unexpected = int(np.count_nonzero(changed & ~allowed))
    assert unexpected == 0, f"Unexpected changed pixels: {unexpected}"
    checks.update(unexpected_changed_pixels=unexpected,
                  accepted_unchanged_regions_pixel_exact=True,
                  standalone_A_D_F_G_byte_identical=True,
                  C_spectra_centroids_and_event_selection_preserved=True,
                  complete_png_rendered_from_delivered_pdf=True)
    make_pptx(output)
    print("Complete PDF/PNG and 8-slide PPTX generated; preservation checks passed", flush=True)

    script_files = [Path(__file__), Path(layout.__file__), Path(shared.__file__)]
    for script in script_files:
        shutil.copy2(script, output / "source" / script.name)
    descriptions = {
        "complete-layout": "按2026-10-19批注将原B拆成B与C，后续面板顺延为D–G。B简化标题并以μV显示波形，E去掉置换示意后与D热图上下对齐。",
        "panela": "保留当前定稿的Y1真实脑投影、A杆和12条双极波形。此独立面板与定稿文件逐字节一致。",
        "panelb": "上图为178段人工标注HFO波形（黑）及均值（黄），从V换算为μV。中图为平均平滑幅度谱的自然对数，下图为相对t≤0.15 s基线归一化后的平均谱；谱值不变，共用0–0.6 s时间轴。",
        "panelc": "原B右侧的Y1 A3–A9三个群体事件独立为C。事件、谱图、完整窗口S³时频质心及±150 ms显示范围均保持；Y1患者标记保留。",
        "paneld": "原C顺延为D，展示18通道、18,190个事件的时间顺序及rank分布。独立文件与原定稿C逐字节一致。",
        "panele": "原D顺延为E，去掉上方Perm/Orig Pats/Mean Pat整段示意，放大MI统计轴。保留20位Yuquan和20位Epilepsiae患者的全部Data/Null点、原小提琴/箱线图、显著性与数值；统计单位为患者。",
        "panelf": "原E顺延为F，保留TA=13,160和TB=5,030的冻结聚类及rank曲线。独立文件与原定稿E逐字节一致。",
        "panelg": "原F顺延为G，保留总体与模板内MI散点和配对inset。独立文件与原定稿F逐字节一致。",
    }
    (output / "figures/README.md").write_text(
        "# Figure 1批注修订：A–G\n\n状态：EXECUTED_PENDING_AUTHOR_VISUAL_REVIEW。\n\n" +
        "\n\n".join(f"### fig1-{key}.png / .pdf\n\n{value}\n\n**关注点**：面板编号、坐标单位及可读性；等待作者目视检查。"
                      for key, value in descriptions.items()) +
        "\n\n### fig1-author-comments.pptx\n\n8页，依次为完整图与A–G独立面板，嵌入本包PNG。矢量内容见各PDF。\n\n**关注点**：用于审阅和排版；PPTX中的面板为图像。\n")
    (output / "README.md").write_text(
        "# Fig1：2026-10-19批注修订\n\n"
        "状态：EXECUTED_PENDING_AUTHOR_VISUAL_REVIEW。来源为作者已接受的y1_final_20261010；当前正式指针保持。\n\n"
        "| 新面板 | 原面板 | 变化 |\n|---|---|---|\n"
        "| A | A | 保留 |\n| B | B左 | 去标题，V→μV，统一时间轴 |\n"
        "| C | B右 | 独立面板，保留Y1 |\n| D | C | 顺延 |\n"
        "| E | D | 去掉上方示意，统计图区放大 |\n| F | E | 顺延 |\n| G | F | 顺延 |\n\n"
        "B的178段事件及两个谱图的定义见figures/README.md，图内不再写Raw、Normalized或HFO n=178。"
        "原D的置换逻辑应由正文/图注说明。本次只交付图件修订，没有改写冻结的论文全文；采用本版时，原C/D/E/F引用需对应更新为D/E/F/G，原B右侧对应C。\n\n"
        "复现入口：scripts/paper_figures/revise_fig1_author_comments.py。完整PNG由同一交付PDF以300 dpi渲染；"
        "validation.json记录数组、统计和未修改区域的核验。运行环境与命令见reproduction.json。\n")
    write_json(output / "reproduction.json", dict(
        command=[sys.executable, "-B", str(Path(__file__).resolve()), "--output-dir", str(output)],
        environment={k: os.environ.get(k, "") for k in ("PYTHONPATH", "LD_LIBRARY_PATH")},
        package_versions={"numpy": np.__version__, "matplotlib": matplotlib.__version__},
        input_package=str(BASE), accepted_version="y1_final_20261010",
        source_hashes={str(path): sha(path) for path in script_files}))
    assert sha(BASE / "current_revision.json") == pointer_hash
    assert verify(BASE) == accepted_report
    write_json(output / "metadata.json", dict(
        version=VERSION, status="EXECUTED_PENDING_AUTHOR_VISUAL_REVIEW",
        human_visual_acceptance="PENDING", source_package=str(BASE),
        panels=list("ABCDEFG"), scientific_data_and_statistics_unchanged=True,
        source_panel_mapping={"A":"A", "B":"B left", "C":"B right", "D":"C", "E":"D", "F":"E", "G":"F"},
        source_C_region_inches=c_box, removal_regions_inches=REMOVE, checks=checks,
        outputs={str(p.relative_to(output)): sha(p) for p in (output / "figures").iterdir()}))
    write_json(output / "validation.json", dict(
        status="PASS", human_visual_acceptance="PENDING", checks=checks,
        accepted_package_verified_before_and_after=accepted_report,
        current_revision_pointer_unchanged=True, pptx_slides=8))
    print(output, flush=True)


if __name__ == "__main__":
    main()
