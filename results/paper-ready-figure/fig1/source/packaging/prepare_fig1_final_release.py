from pathlib import Path
import ast, csv, hashlib, json, shutil, sys
import numpy as np

ROOT=Path('/home/honglab/leijiaxin/HFOsp')
CANON=ROOT/'results/paper-ready-figure/fig1'
ACCEPTED=CANON/'revisions/y1_tighter_lower_row_20261010'
LAYOUT_INPUT=CANON/'revisions/y1_a_left_label_clear_lead_20261010'
STAGE=Path('/tmp/fig1-final-staging-20261010')
ARCHIVE=ROOT/'results/paper-ready-figure/archive/2026-10-10_fig1_finalization/fig1'
assert not STAGE.exists()
STAGE.mkdir()
sys.path[:0]=[str(ROOT),str(ROOT/'scripts/paper_figures'),'/tmp/fig1_release_ppt_deps']
from scripts.paper_figures import plot_fig1_interictal_hfo_temporal_scaffold as original
from scripts.paper_figures import restore_fig1_legacy_spectrum as spectrum
from scripts.paper_figures.build_fig1_current import digest

def write(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n')

copied=[]
def save(src,rel):
    dst=STAGE/rel;dst.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(src,dst)
    assert digest(src)==digest(dst)
    copied.append(dict(source=str(src),backup=str(rel),sha256=digest(dst),bytes=dst.stat().st_size))
    return dst

for p in (ACCEPTED/'figures').iterdir():
    if p.suffix in ('.png','.pdf'):
        target=('quality_control/figures' if 'full-window' in p.name else 'figures')
        save(p,Path(target)/p.name)
save(ACCEPTED/'spectrum_contract.json','spectrum_contract.json')
save(ACCEPTED/'metadata.json','source/accepted_revision_metadata.json')
save(ACCEPTED/'validation.json','source/accepted_revision_validation.json')
save(CANON/'current_revision.json','source/accepted_revision_pointer.json')
# The exact last compositor input is kept separate from scientific inputs.
for p in LAYOUT_INPUT.rglob('*'):
    rel=p.relative_to(LAYOUT_INPUT)
    if p.is_file() and (rel.parts[0]=='figures' or str(rel) in ('metadata.json','spectrum_contract.json')):
        save(p,Path('source/layout_input')/rel)

a=CANON/'revisions/y1_a7_zoom_final_20261009/source/panel_a_snapshot'
for name in ('selection.json','brain_projection.json','recording_and_geometry.npz','original_recording_and_geometry.npz','y1_brain.png'):
    save(a/'source'/name,Path('data/a')/name)
save(a/'metadata.json','data/a/geometry_metadata.json')
local=CANON/'candidates/y1_local_rank_peak_profiles_20261010'
for name in ('display_ranks_18.npz','rank_histograms.npz'):
    save(local/'source'/name,Path('data/ce')/name)
for name in ('display_channel_selection.csv','display_channel_selection.json','display_rank_summary.json'):
    save(local/name,Path('data/ce')/name)
save(local/'source/hfo_showcase.npz','data/b/hfo_showcase.npz')
for record in ('FA134AX6','FA134AXF'):
    save(spectrum.OUT/f'source/{record}_segment.npz',Path('data/b')/f'{record}_segment.npz')
contract=json.loads((ACCEPTED/'spectrum_contract.json').read_text())
events=spectrum.display_events(contract,write_metadata=False,event_selection=[('FA134AX6',1559),('FA134AX6',1562),('FA134AXF',1494)])
for i,(event,info) in enumerate(zip(events,contract['display_events'])):
    np.savez_compressed(STAGE/f'data/b/event_{i+1}_{info["record"]}_{info["event_index"]}.npz',**event)
for record in ('FA134AX6','FA134AXF'):
    with np.load(STAGE/f'data/b/{record}_segment.npz') as z:
        recomputed=spectrum.legacy_spectrum(z['signals_stitched_V'],z['split_borders_sec'])
        for key,value in zip(('specs','times','freqs','centers','unnormalized_specs'),recomputed):
            np.testing.assert_array_equal(z[key],value)

records=original._load_temporal_records(); original._assert_masked_mi_records(records)
assert len(records)==40
record=next(r for r in records if r['dataset']=='yuquan' and r['subject']=='zhangkexuan')
arrays=original._load_exemplar_arrays(record,10**9)
meta=json.loads((ACCEPTED/'metadata.json').read_text())
for key,expected in meta['original_array_hashes'].items():
    assert hashlib.sha256(np.ascontiguousarray(arrays[key]).tobytes()).hexdigest()==expected,key
np.savez_compressed(STAGE/'data/ce/original_arrays_26.npz',**{k:np.asarray(v) for k,v in arrays.items() if isinstance(v,(np.ndarray,list))})
save(Path(meta['source_record']),'data/ce/subject_record.json')
reduced=[]
for r in records:
    path=original.MASKED_ROOT/f'per_subject/{r["dataset"]}_{r["subject"]}.json'
    reduced.append(dict(dataset=r['dataset'],subject=r['subject'],legacy_mi=r['legacy_mi'],
        adaptive_cluster={k:r['adaptive_cluster'][k] for k in ('overall_tau','within_cluster_tau_mean','chosen_k')},
        source_path=str(path),source_sha256=digest(path)))
write(STAGE/'data/df/cohort_plot_records.json',reduced)
with (STAGE/'data/df/cohort_plot_values.csv').open('w') as f:
    fields=['dataset','subject','masked_mi','permutation_null_median','overall_mi','within_template_mi']
    writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader()
    for r in reduced:
        writer.writerow(dict(dataset=r['dataset'],subject=r['subject'],masked_mi=r['legacy_mi']['mi_mean'],permutation_null_median=r['legacy_mi']['permuted_mean_median'],overall_mi=r['adaptive_cluster']['overall_tau'],within_template_mi=r['adaptive_cluster']['within_cluster_tau_mean']))
raw=Path('/mnt/yuquan_data/yuquan_24h_edf/zhangkexuan')
for record in ('FA134AX4','FA134AX6','FA134AXF'):
    for suffix in ('_lagPat.npz','_packedTimes.npy'):
        save(raw/(record+suffix),Path('data/raw_artifact_excerpt')/(record+suffix))

# Freeze all Figure 1 producer revisions plus their local Python import closure.
seeds=set((ROOT/'scripts/paper_figures').glob('*fig1*.py'))
seeds.update(ROOT/p for p in ('scripts/paper_figures/build_main_figures_1_2.py','src/lagpat_rank_audit.py','src/paper_figure_typography.py'))
queue=list(seeds);visited=set()
def module_file(name,parent):
    candidates=[ROOT/Path(*name.split('.')).with_suffix('.py'),parent/Path(*name.split('.')).with_suffix('.py'),ROOT/'scripts'/Path(*name.split('.')).with_suffix('.py'),ROOT/'scripts/paper_figures'/Path(*name.split('.')).with_suffix('.py')]
    return next((p for p in candidates if p.is_file()),None)
while queue:
    p=queue.pop()
    if p in visited:continue
    visited.add(p)
    try:tree=ast.parse(p.read_text())
    except (SyntaxError,UnicodeDecodeError):continue
    modules=[]
    for node in ast.walk(tree):
        if isinstance(node,ast.Import):modules.extend(a.name for a in node.names)
        if isinstance(node,ast.ImportFrom) and node.module:
            modules.append(node.module)
            modules.extend(node.module+'.'+a.name for a in node.names if a.name!='*')
    for name in modules:
        child=module_file(name,p.parent)
        if child is not None and child not in visited:queue.append(child)
for p in sorted(visited):save(p,Path('source/code_snapshot')/p.relative_to(ROOT))
for p in (spectrum.ORIGINAL,Path(meta['hfo_source']['legacy_code_reference'])):
    save(p,Path('source/code_snapshot')/p.relative_to(ROOT))
save(ROOT/'docs/figure_style_guide.md','source/figure_style_guide_at_acceptance.md')

panel_contracts={
 'A':{'producer':'build_fig1a_recording_chain.py + build_fig1_y1_local_rank.py + revise_fig1_a_left_label_clear_lead.py','data':'data/a','meaning':'Y1 brain projection, 12 A-shaft bipolar waveforms, A7/A9 colors; 80–250 Hz; 3 x 0.16 s'},
 'B':{'producer':'restore_fig1_legacy_spectrum.py + revise_fig1_b_compact_third_event.py + revise_fig1_b_visual_alignment.py','data':'data/b','meaning':'178 real HFO snippets; original full-window S^3 centroids for FA134AX6/1559,1562 and FA134AXF/1494; +/-150 ms display'},
 'C/E':{'producer':'build_fig1_y1_local_rank.py','data':'data/ce','meaning':'18 display channels, 18,190 events; frozen TA=13,160 and TB=5,030; display ranks 1–18, nonparticipants blank; original 26-channel arrays retained'},
 'D/F':{'producer':'plot_fig1_interictal_hfo_temporal_scaffold.py + build_fig1_readability_review.py','data':'data/df','meaning':'20 Yuquan + 20 Epilepsiae; masked MI vs null and overall/within-template paired statistics unchanged'},
}
write(STAGE/'source/panel_data_contract.json',panel_contracts)
write(STAGE/'data/backup_manifest.json',dict(scope='Partial scientific backup sufficient to retain exact figure inputs; not a full EDF/anatomical reconstruction dataset backup.',copied_files=copied,generated=['ce/original_arrays_26.npz','df/cohort_plot_records.json','df/cohort_plot_values.csv','b/event_1_FA134AX6_1559.npz','b/event_2_FA134AX6_1562.npz','b/event_3_FA134AXF_1494.npz'],raw_root=str(raw),no_raw_data_modified=True))

from pptx import Presentation
from pptx.util import Inches
from PIL import Image
prs=Presentation();prs.slide_width=Inches(19.5);prs.slide_height=Inches(15.65)
for name in ('complete-layout',*[f'panel{x}' for x in 'abcdef']):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);p=STAGE/f'figures/fig1-{name}.png'
    with Image.open(p) as im: w,h=im.size
    scale=min(prs.slide_width/w,prs.slide_height/h)
    pw,ph=int(w*scale),int(h*scale)
    slide.shapes.add_picture(str(p),(prs.slide_width-pw)//2,(prs.slide_height-ph)//2,width=pw,height=ph)
    slide.notes_slide.notes_text_frame.text=f'Figure 1 {name}; author accepted 2026-10-10. Frozen scientific image, no resampling in source asset. Editable vector version supplied separately as PDF.'
prs.save(STAGE/'figures/fig1-final.pptx')

meta.update(status='AUTHOR_ACCEPTED_FINAL',human_visual_acceptance='ACCEPTED_2026-10-10',panel_a_human_visual_acceptance='ACCEPTED_2026-10-10',version='y1_final_20261010',
    producer='scripts/paper_figures/build_fig1_current.py',
    source_revision=str(ARCHIVE/'revisions/y1_tighter_lower_row_20261010'),
    source_record_external=meta['source_record'],source_record='data/ce/subject_record.json',
    source_manifest='source/panel_data_contract.json',data_backup='data/backup_manifest.json',
    spectrum_contract='spectrum_contract.json',validation='validation.json',
    changed_panels=[],preservation={'promoted_files_byte_identical_to_author_accepted_revision':True})
meta['outputs']={str(p.relative_to(STAGE)):digest(p) for p in (STAGE/'figures').iterdir() if p.is_file()}
write(STAGE/'metadata.json',meta)
write(STAGE/'figure1_panel_registry.json',dict(status='AUTHOR_ACCEPTED_FINAL',version='y1_final_20261010',producer='scripts/paper_figures/build_fig1_current.py',source_manifest='source/panel_data_contract.json',data_backup='data/backup_manifest.json',panels={x:[f'figures/fig1-panel{x.lower()}.{e}' for e in ('png','pdf')] for x in 'ABCDEF'},complete=[f'figures/fig1-complete-layout.{e}' for e in ('png','pdf')],ppt='figures/fig1-final.pptx',panel_letters_in_individual_files=False))
write(STAGE/'current_revision.json',dict(status='AUTHOR_ACCEPTED_FINAL',version='y1_final_20261010',selected_patient='Y1',panels=list('ABCDEF'),revision_root=str(CANON),package_relative='results/paper-ready-figure/fig1',producer='scripts/paper_figures/build_fig1_current.py',metadata='results/paper-ready-figure/fig1/metadata.json',complete_figure='results/paper-ready-figure/fig1/figures/fig1-complete-layout.pdf',complete_figure_png='results/paper-ready-figure/fig1/figures/fig1-complete-layout.png',ppt='results/paper-ready-figure/fig1/figures/fig1-final.pptx',human_visual_acceptance='ACCEPTED_2026-10-10',panel_a_human_visual_acceptance='ACCEPTED_2026-10-10',previous_versions_archive=str(ARCHIVE),source_manifest='results/paper-ready-figure/fig1/source/panel_data_contract.json',data_backup='results/paper-ready-figure/fig1/data/backup_manifest.json',data_and_statistics_unchanged=True))
write(STAGE/'validation.json',dict(status='PASS',human_visual_acceptance='ACCEPTED_2026-10-10',all_14_primary_png_pdf_files_byte_identical=True,spectrum_arrays_exact_against_original_algorithm=True,original_26_channel_hashes_match=True,events=18190,cohort_n=40,ppt_slides=7,independent_panels='A–F',PNG_PDF_visual_self_review='accepted revision previously inspected; publish unchanged'))

text='''# 当前 Figure 1：作者定稿，2026-10-10

唯一正式入口为本目录的 `current_revision.json`；完整图与A–F单独面板在 `figures/`。`figures/fig1-final.pptx`含7页：完整图及六个单独面板。所有PNG/PDF与作者接受的紧凑行距版本逐字节一致。

## 代码与数据

当前入口：`scripts/paper_figures/build_fig1_current.py`。原生绘图代码及本次全部修订代码的冻结副本在 `source/code_snapshot/`，面板与代码/数据的对应关系在 `source/panel_data_contract.json`。

`data/a`包含真实波形、触点坐标及脑表面投影；`data/b`包含178段HFO及完整打包事件谱、质心和显示例；`data/ce`同时备份原26通道数组与18通道显示派生数组、冻结分类及昼夜标记；`data/df`为40位患者原统计读出。原始EDF和完整解剖重建仍在原数据盘，本包是部分备份，未改动原始数据。

## 校验与复现

运行 `python scripts/paper_figures/build_fig1_current.py` 校验当前包；加 `--output-dir /tmp/fig1-rebuild` 可在空目录复现定稿拼版，验证完整PNG及单面板一致。需要 numpy、pillow、pypdf、matplotlib；原始科学producer另依赖项目环境。冻结PDF/PNG图层负责精确版式复现；如果修改数据或科学计算，应使用对应面板的原producer及数据合同，不能将拼版复现表述为重新运行原始检测。

## 归档与协作

旧正式图、所有旧候选和修订已迁入 `../archive/2026-10-10_fig1_finalization/fig1/`。旧路径仅作历史来源，不再是绘图入口。主目录 `docs/current_figure1.md`、登记表和各工作树AGENTS入口共同指向本版；任何后续修改从本版派生候选，不直接覆盖已接受包。
'''
(STAGE/'README.md').write_text(text)
(STAGE/'figures/README.md').write_text('# Figure 1 正式定稿图\n\n'+
 '\n\n'.join(f'### fig1-{name}.png / .pdf\n\n{desc}\n\n**关注点**：作者于2026-10-10接受，本次仅提升为正式入口，图像内容保持。' for name,desc in [
 ('complete-layout','完整A–F图，E/F整行上移6.35 mm后的定稿布局。'),
 ('panela','Y1真实脑定位、A杆局部放大及12条双极波形。Y1左对齐，A7紫色和A9蓝色对应。'),
 ('panelb','178段HFO谱及Y1的三个群体事件。使用原完整事件S³质心和统一±150 ms显示窗。'),
 ('panelc','Y1的18通道时间顺序热图、Day/Night及rank分布。全部18,190事件保留。'),
 ('paneld','原示意及40位患者的masked MI与置换null统计。标题已分开。'),
 ('panele','冻结TA/TB重排热图和rank均值及标准差。原始26通道数据不删减。'),
 ('panelf','40位患者总体与模板内MI及配对inset。原统计语法保留。')])+
 '\n\n### fig1-final.pptx\n\n含完整图和A–F共7页，嵌入本版原PNG。矢量图另见各PDF。\n\n**关注点**：各页图像比例保持；PPT内图片不声称为可编辑原生曲线。\n')
(STAGE/'quality_control/figures/README.md').write_text('### fig1-spectrum-full-window-check.png / .pdf\n\n展示三个完整500 ms计算窗口，虚线表示主图±150 ms显示范围。质心仍由完整窗口计算。\n\n**关注点**：此图供方法核对，不属于主图A–F。\n')
manifest=dict(status='AUTHOR_ACCEPTED_FINAL',version='y1_final_20261010',accepted_on='2026-10-10',files={str(p.relative_to(STAGE)):dict(sha256=digest(p),bytes=p.stat().st_size) for p in sorted(STAGE.rglob('*')) if p.is_file()})
write(STAGE/'release_manifest.json',manifest)
print('STAGED',STAGE,'files',len(manifest['files']),'MB',round(sum(x['bytes'] for x in manifest['files'].values())/1e6,2),flush=True)
